from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "data" / "Sola_Stavanger_Sandnes_schools_map_ready.xlsx"
OUT = ROOT / "workbook-preview"
FONT_PATH = Path("C:/Windows/Fonts/arial.ttf")
BOLD_PATH = Path("C:/Windows/Fonts/arialbd.ttf")


def font(size: int, bold: bool = False):
    path = BOLD_PATH if bold and BOLD_PATH.exists() else FONT_PATH
    return ImageFont.truetype(str(path), size=size) if path.exists() else ImageFont.load_default()


def render_table(sheet, path: Path, rows: list[int], columns: int) -> None:
    widths = []
    for column in range(1, columns + 1):
        letter = sheet.cell(1, column).column_letter
        excel_width = sheet.column_dimensions[letter].width or 10
        widths.append(max(80, min(260, int(excel_width * 8))))
    heights = [54 if row == 6 else 34 for row in rows]
    image = Image.new("RGB", (sum(widths) + 2, sum(heights) + 2), "white")
    draw = ImageDraw.Draw(image)
    y = 1
    for row, height in zip(rows, heights):
        x = 1
        for column, width in enumerate(widths, start=1):
            cell = sheet.cell(row, column)
            fill = cell.fill.fgColor.rgb
            fill_color = f"#{fill[-6:]}" if fill and fill not in {"00000000", "000000"} else ("#17324D" if row == 6 else "#FFFFFF")
            draw.rectangle((x, y, x + width, y + height), fill=fill_color, outline="#CBD5E1", width=1)
            value = "" if cell.value is None else str(cell.value)
            if len(value) > 32:
                value = value[:29] + "…"
            color = "#FFFFFF" if row == 6 else "#17202A"
            draw.text((x + 5, y + 7), value, fill=color, font=font(12, bold=row == 6))
            x += width
        y += height
    image.save(path)


def render_map_sheet(sheet, path: Path) -> None:
    image = Image.new("RGB", (1500, 720), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), str(sheet["A2"].value), fill="#17324D", font=font(30, bold=True))
    y = 135
    for label_row in (4, 6, 8, 10, 11, 13):
        label = sheet.cell(label_row, 1).value or ""
        text = sheet.cell(label_row, 2).value or "" if label_row != 13 else sheet.cell(label_row, 1).value or ""
        if label_row == 13:
            draw.text((70, y), str(text), fill="#566575", font=font(16))
            break
        draw.rounded_rectangle((70, y, 1430, y + 72), radius=10, fill="#FFFFFF", outline="#DCE5ED")
        draw.text((95, y + 18), str(label), fill="#17324D", font=font(17, bold=True))
        draw.text((290, y + 18), str(text), fill="#263746", font=font(16))
        y += 92
    image.save(path)


OUT.mkdir(parents=True, exist_ok=True)
workbook = load_workbook(BOOK, data_only=False)
all_schools = workbook["All schools"]
visible_rows = [6] + [row for row in range(7, all_schools.max_row + 1) if not all_schools.row_dimensions[row].hidden][:22]
render_table(all_schools, OUT / "all-schools.png", visible_rows, all_schools.max_column)
render_map_sheet(workbook["Map"], OUT / "map-sheet.png")
print(OUT)

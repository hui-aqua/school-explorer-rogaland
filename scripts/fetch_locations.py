from __future__ import annotations

import html
import json
import re
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_XLSX = ROOT / "data" / "original" / "Sola_Stavanger_Sandnes_schools.xlsx"
CACHE_FILE = ROOT / "data" / "locations.json"


def load_rows() -> list[dict]:
    workbook = load_workbook(SOURCE_XLSX, data_only=True, read_only=False)
    sheet = workbook["All schools"]
    headers = [cell.value for cell in sheet[6]]
    rows = []
    for row_number in range(7, sheet.max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, sheet.max_column + 1)]
        record = dict(zip(headers, values))
        record["_row"] = row_number
        record["_visible"] = not bool(sheet.row_dimensions[row_number].hidden)
        rows.append(record)
    return rows


def parse_location(page: str) -> dict:
    coords = re.search(r"[?&]mlat=([-0-9.]+)&(?:amp;)?mlon=([-0-9.]+)", page)
    if not coords:
        coords = re.search(r"marker=([-0-9.]+)%2C([-0-9.]+)", page)
    address_match = re.search(r"Skolen ligger i\s+(.+?)\.</p>", page, flags=re.I | re.S)
    if not address_match:
        address_match = re.search(r"Skolen ligger i\s+(.+?)\.", page, flags=re.I | re.S)
    address = None
    if address_match:
        address = re.sub(r"<[^>]+>", " ", address_match.group(1))
        address = html.unescape(re.sub(r"\s+", " ", address)).strip()
    return {
        "latitude": float(coords.group(1)) if coords else None,
        "longitude": float(coords.group(2)) if coords else None,
        "address": address,
    }


def fetch_one(record: dict) -> tuple[str, dict]:
    url = record.get("Source URL")
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "school-explorer-rogaland/1.0 (+data provenance project)"},
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        page = response.read().decode("utf-8", errors="replace")
    location = parse_location(page)
    location.update(
        {
            "school": record.get("School"),
            "municipality": record.get("Municipality"),
            "organisation_number": str(record.get("Organisation number") or ""),
            "source_url": url,
        }
    )
    return str(record.get("Organisation number") or ""), location


def main() -> None:
    rows = load_rows()
    existing = {}
    if CACHE_FILE.exists():
        existing = json.loads(CACHE_FILE.read_text(encoding="utf-8"))

    pending = [row for row in rows if str(row.get("Organisation number") or "") not in existing]
    failures: list[dict] = []
    if pending:
        with ThreadPoolExecutor(max_workers=6) as executor:
            futures = {executor.submit(fetch_one, row): row for row in pending}
            for future in as_completed(futures):
                row = futures[future]
                try:
                    key, location = future.result()
                    existing[key] = location
                except Exception as exc:  # keep partial cache and report exact failures
                    failures.append({"school": row.get("School"), "url": row.get("Source URL"), "error": str(exc)})
                time.sleep(0.03)

    CACHE_FILE.write_text(
        json.dumps(existing, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    geocoded = sum(1 for value in existing.values() if value.get("latitude") is not None)
    print(json.dumps({"schools": len(rows), "geocoded": geocoded, "failures": failures}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "original" / "Sola_Stavanger_Sandnes_schools.xlsx"
HTML = ROOT / "legacy" / "school-map.html"


def workbook_report() -> dict:
    wb = load_workbook(XLSX, data_only=False)
    report: dict = {"sheets": []}
    for ws in wb.worksheets:
        hidden_rows = [idx for idx, dim in ws.row_dimensions.items() if dim.hidden]
        table_ranges = {
            name: (table.ref if hasattr(table, "ref") else str(table))
            for name, table in ws.tables.items()
        }
        values = list(
            ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, 8),
                min_col=1,
                max_col=min(ws.max_column, 24),
                values_only=True,
            )
        )
        report["sheets"].append(
            {
                "title": ws.title,
                "dimensions": ws.calculate_dimension(),
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "auto_filter": ws.auto_filter.ref,
                "hidden_rows": hidden_rows[:30],
                "hidden_row_count": len(hidden_rows),
                "tables": table_ranges,
                "sample": values,
            }
        )
    return report


def html_report() -> dict:
    text = HTML.read_text(encoding="utf-8", errors="replace")
    scripts = re.findall(r"<script\b[^>]*?(?:src=[\"']([^\"']+)[\"'])?[^>]*>", text, re.I)
    styles = re.findall(r"<link\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>", text, re.I)
    title = re.search(r"<title>(.*?)</title>", text, re.I | re.S)
    snippets = {}
    for token in ["L.map", "const schools", "var schools", "schoolData", "xlsx", "SheetJS", "fetch("]:
        pos = text.find(token)
        if pos >= 0:
            snippets[token] = text[max(0, pos - 180) : pos + 420]
    return {
        "bytes": HTML.stat().st_size,
        "characters": len(text),
        "title": title.group(1).strip() if title else None,
        "script_srcs": [item for item in scripts if item],
        "link_hrefs": styles,
        "inline_script_count": len(re.findall(r"<script\b(?![^>]*src=)", text, re.I)),
        "snippets": snippets,
    }


print(json.dumps({"workbook": workbook_report(), "html": html_report()}, ensure_ascii=False, indent=2, default=str))
